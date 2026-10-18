"""Fill the AI Builder Cup template without changing its type.

Starts from the official template. Does not delete slides, does not restyle
titles, and does not introduce Arial. New text uses the embedded faces:

  Google Sans Flex SemiBold   headings we add
  Google Sans Flex Medium     already on the template titles
  Google Sans Flex            body, same as "Team name:" on slide 1

The theme default is Arial. Every new run sets latin/ea/cs/sym or PowerPoint
falls back to that default.
"""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC = "Template/Submission Template _ AI Builder Cup.pptx"
OUT = "Template/Out/Multimodal Studio - AI Builder Cup.pptx"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
INK = RGBColor(0x1F, 0x1F, 0x1F)
SOFT = RGBColor(0x3C, 0x40, 0x43)

SEMI = "Google Sans Flex SemiBold"
MED = "Google Sans Flex Medium"
REG = "Google Sans Flex"


def _face(run, name: str, size_pt: float) -> None:
    """Pin every script slot. A latin-only set still renders as Arial."""
    run.font.size = Pt(size_pt)
    run.font.bold = False
    run.font.italic = False
    run.font.color.rgb = INK
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs", "sym"):
        el = rPr.find(f"{A}{tag}")
        if el is None:
            el = etree.SubElement(rPr, f"{A}{tag}")
        el.set("typeface", name)


def _box(slide, left, top, width, height):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    return tf


def write(slide, lines, *, left=0.32, top=1.5, width=9.35, height=3.85, size=14, font=REG):
    """lines: str, or (text, font, size)."""
    tf = _box(slide, left, top, width, height)
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, face, sz = item
        else:
            text, face, sz = item, font, size
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = text
        _face(run, face, sz)
    return tf


def find(prs, needle: str):
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            blob = "\n".join(p.text for p in sh.text_frame.paragraphs)
            if needle in blob:
                return slide, sh
    raise KeyError(needle)


def fill_team(prs) -> None:
    """Keep the template's own runs and faces. Only fill the blanks it left."""
    _, shape = find(prs, "Team name:")
    tf = shape.text_frame
    # 0 Team Details (SemiBold 21), 2 name, 3 leader, 4 problem — all Google Sans Flex 18.
    tf.paragraphs[2].runs[0].text = "Team name: Multimodal Studio"
    tf.paragraphs[3].runs[0].text = (
        "Team leader: Anil Kumar  ·  anil8860556664@gmail.com"
    )
    tf.paragraphs[4].runs[0].text = (
        "Member: M Sai Teja  ·  macharlasaiteja@gmail.com.  "
        "A chapter still becomes slides by hand, then a recording. "
        "Students get a flat read-aloud. We wanted a video they would actually replay."
    )


def steps(slide) -> None:
    """The process-flow slide asks for a diagram. Six plain steps, template type."""
    labels = ["Drop a file", "Read it", "Write slides", "Speak them", "Film frames", "MP4"]
    left0 = 0.32
    gap = 0.12
    w = 1.48
    top = 2.15
    h = 0.95
    for i, label in enumerate(labels):
        x = left0 + i * (w + gap)
        sh = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(w), Inches(h)
        )
        sh.adjustments[0] = 0.15
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
        sh.line.color.rgb = RGBColor(0xDA, 0xDC, 0xE0)
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        try:
            tf._txBody.bodyPr.set("anchor", "ctr")
        except Exception:
            pass
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        _face(run, SEMI, 13)
    write(
        slide,
        [
            "Same path if you uploaded the file, or if the overnight job picked a paper.",
            "Speaking is Gemini 3.8 Live, so the voice has pauses and a bit of warmth. It is not a flat read-out.",
        ],
        top=3.35,
        height=1.7,
        size=14,
    )


def shots(slide) -> None:
    write(
        slide,
        ["These are screens from the app we run, not drawings of a future one."],
        top=1.42,
        height=0.4,
        size=13,
    )
    pics = [
        ("docs/screenshots/studio.png", "Where a run starts. File on the left, format and voice, then run."),
        ("docs/screenshots/voices.png", "Kore is the voice in the demo. You can hear it before you commit."),
        ("docs/screenshots/history.png", "Old runs stay here. Open one again instead of starting over."),
    ]
    x = 0.32
    for img, cap in pics:
        slide.shapes.add_picture(img, Inches(x), Inches(1.95), Inches(3.0), Inches(1.49))
        write(slide, [cap], left=x, top=3.52, width=3.0, height=1.5, size=12)
        x += 3.15


def main() -> None:
    prs = Presentation(SRC)

    fill_team(prs)

    brief, _ = find(prs, "Brief about the idea")
    write(
        brief,
        [
            "Anil Kumar leads the team. M Sai Teja builds with him. We kept turning chapters into slides and then sitting down to record.",
            "Multimodal Studio is for a teacher, a tutor, or a student who wants the idea to stick. Drop the PDF or the notes, say the age, and you get a video with a voice already on it.",
            "The fun is in the timing. The highlight moves when the sentence does, so it feels like someone explaining, not a page being read out. Gemini writes it. Gemini 3.8 Live speaks it. The review link opens with no login.",
        ],
    )

    opp, _ = find(prs, "How different is it")
    write(
        opp,
        [
            ("How it is different", SEMI, 14),
            "Slide tools stop at the deck. Read-aloud tools stay in one tone and ignore the picture. A teacher was still doing the joining. This goes from the file to a video a class can replay.",
            ("How it solves it", SEMI, 14),
            "The slow part was the middle: slides a student can follow, a voice with some life in it, and a highlight that walks with the sentence. That middle is the product.",
            ("What we would defend", SEMI, 14),
            "One Gemini key. A voice that does not sound like a station announcement. A link anyone on the review team can open.",
        ],
        top=2.32,
        height=3.05,
        size=13,
    )

    feats, _ = find(prs, "List of features")
    write(
        feats,
        [
            "A teacher can turn one chapter into a YouTube lesson, a Short, or a Reel. The length follows the format.",
            "Age and how much the viewer already knows change the wording. A class of 16-year-olds does not get the same script as a beginner.",
            "Gemini 3.8 Live speaks it. Kore is the demo voice. You can preview before you run.",
            "A long PDF can be split into topics, then one video or a few, so a unit does not have to be one 40-minute file.",
            "It flags a slide that repeats the one before it, and the on-screen highlight stays with the voice.",
            "Title, description, and thumbnail come out with the file.",
        ],
        size=13,
    )

    flow, _ = find(prs, "Process flow")
    write(
        flow,
        ["Two ways in. You drop something, or the admin queue picks a paper overnight. After that it is one path."],
        top=1.42,
        height=0.6,
        size=13,
    )
    steps(flow)

    wire, _ = find(prs, "Wireframes")
    wire.shapes.add_picture(
        "docs/screenshots/studio.png", Inches(0.32), Inches(1.5), Inches(6.15), Inches(3.05)
    )
    write(
        wire,
        [
            ("What you actually see", SEMI, 13),
            "The file goes on the left.",
            "Format and voice are next to it.",
            "You press run and wait on the same page.",
            "The MP4 plays there when it is done. History keeps the old ones.",
        ],
        left=6.65,
        top=1.5,
        width=3.05,
        height=3.4,
        size=13,
    )

    arch, _ = find(prs, "Architecture diagram")
    write(
        arch,
        [
            "Browser  →  our Python app on Cloud Run  →  Gemini  →  an MP4",
            "The page is the Studio you saw. The server asks Gemini 3.8 Flash (or Argon) to read the file and write the slides. Gemini 3.8 Live speaks them. A headless Chrome photographs the slides. FFmpeg puts the sound on.",
            "Each job is saved on the machine and copied to Firestore. Firebase Hosting only serves this pitch. The app itself is Cloud Run.",
        ],
        size=14,
    )

    tech, _ = find(prs, "Technologies to be used")
    write(
        tech,
        [
            "Writing, and reading pages: Gemini 3.8 Flash. Gemini 4 Argon when the source is dense and we want it to think harder.",
            "Voice: Gemini 3.8 Live. Same family as the in-app assistant. We stopped using a separate flat text-to-speech voice.",
            "Thumbnail only: Nano Banana. It does not draw the slides.",
            "App: Python, FastAPI, HTML slides, headless Chrome, FFmpeg.",
            "Hosting: Cloud Run. This page: Firebase Hosting. Job copy: Firestore.",
        ],
        size=14,
    )

    cost, _ = find(prs, "Estimated implementation cost")
    write(
        cost,
        [
            "No GPU. The only bill is the Gemini API.",
            "On Flash, a normal video is a few rupees. Argon costs more, so we turn it on when the source needs it, not for every run.",
            "Cloud Run and Firestore should stay inside the free tier for the number of people who will open this during review.",
            "We have not switched on a paid hosting plan. The API key is the thing that has to be set.",
        ],
    )

    snap, _ = find(prs, "Snapshots of the prototype")
    shots(snap)

    perf, _ = find(prs, "Prototype Performance")
    write(
        perf,
        [
            "There is no benchmark table on purpose. We did not want a number on this slide that we had not measured.",
            "Every run does write down how long each step took. That lives with the job, not in a slide we typed.",
            "Filming is the slow step. We shoot the movement at the start of a slide and hold the frame, so a long voiceover is not a screenshot every fraction of a second.",
            "If Gemini is busy, that step waits and tries again. The run does not stop in the middle and leave a broken file.",
        ],
    )

    later, _ = find(prs, "Additional Details")
    write(
        later,
        [
            "Next we want the Studio to feel more obvious: fewer choices, a clearer next step, and the look of the slides kept in step with the style you picked.",
            "Handwritten teaching templates, separate from the board we have now, so a maths or science idea can look drawn rather than typed.",
            "Manim-style motion for formulas. Predefined YouTuber looks, plus a custom one. Better animation models, both the ones we ship and ones a teacher can add.",
            "New animation skills you can drop in, the way a style is picked today. A Short cut from the long lesson without doing it by hand.",
        ],
    )

    links, shape = find(prs, "GitHub Public Repository")
    # The template box is only tall enough for the three labels. Grow it.
    shape.height = Inches(4.2)
    tf = shape.text_frame
    notes = [
        "GitHub: https://github.com/saitej123/Gemini-Multimodal-Studio",
        "Demo video: 3 to 4 minutes of one real file becoming an MP4. Link still to be pasted.",
        "Product: https://multimodal-studio-469528918286.us-central1.run.app",
        "That link opens in Studio. There is no login.",
    ]
    for text in notes:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(8)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = text
        _face(run, REG, 13)

    # Slides 2 (download note), 15 (blank), and 16 (closing art) stay as the template drew them.
    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
